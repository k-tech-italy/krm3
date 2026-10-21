from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from bs4 import BeautifulSoup
from django.urls import reverse
from testutils.date_utils import _dt
from testutils.factories import ContractFactory, DayEntryFactory, SpecialLeaveReasonFactory
from testutils.permissions import add_permissions
from krm3.core.models import Contract


def _get_resource_table(response, resource):
    soup = BeautifulSoup(response.content, 'html.parser')
    for table in soup.find_all('table', class_='report-table'):
        header = table.find('td', class_='row-header')
        if header and resource.first_name in header.get_text() and resource.last_name in header.get_text():
            return table
    return None


def _get_report_rows(table):
    return {
        row.find('td', class_='row-header').get_text(strip=True): [
            cell.get_text(strip=True) for cell in row.find_all('td', class_='cell-data')
        ]
        for row in table.find_all('tr', class_='row-data')
        if row.find('td', class_='row-header')
    }


def test_timesheet_report_shows_day_entry_hours_and_totals(admin_client):
    contract = ContractFactory()
    resource = contract.resource
    DayEntryFactory(
        contract=contract,
        resource=resource,
        day=_dt('2025-06-05'),
        due_hours=Decimal(8),
        day_hours=Decimal(6),
        travel_hours=Decimal(2),
    )

    response = admin_client.get(reverse('report-month', args=['202506']))

    assert response.status_code == 200
    table = _get_resource_table(response, resource)
    assert table is not None
    rows = _get_report_rows(table)
    assert rows['Regular hours'][0] == '8'
    assert rows['Regular hours'][5] == '8'
    assert rows['Day shift hours'][0] == '6'
    assert rows['Travel'][0] == '2'


@pytest.mark.parametrize(
    'excluded_contract_type',
    [
        Contract.ContractType.CONTRACTOR,
        Contract.ContractType.OTHER,
    ],
)
def test_timesheet_report_only_shows_employees_to_privileged_users(
    admin_client,
    excluded_contract_type,
):
    employee_contract = ContractFactory(
        contract_type=Contract.ContractType.EMPLOYEE,
    )
    excluded_contract = ContractFactory(
        contract_type=excluded_contract_type,
    )

    response = admin_client.get(reverse('report-month', args=['202506']))

    assert response.status_code == 200
    assert _get_resource_table(response, employee_contract.resource) is not None
    assert _get_resource_table(response, excluded_contract.resource) is None


def test_timesheet_report_shows_sick_protocol_and_special_leave(admin_client):
    contract = ContractFactory()
    resource = contract.resource
    reason = SpecialLeaveReasonFactory(title='Parental leave')
    DayEntryFactory(
        contract=contract,
        resource=resource,
        day=_dt('2025-06-09'),
        due_hours=Decimal(8),
        is_sick=True,
        protocol_number='SICK-123',
    )
    DayEntryFactory(
        contract=contract,
        resource=resource,
        day=_dt('2025-06-10'),
        due_hours=Decimal(8),
        special_leave_hours=Decimal(4),
        special_leave_reason=reason,
    )

    response = admin_client.get(reverse('report-month', args=['202506']))

    assert response.status_code == 200
    table = _get_resource_table(response, resource)
    assert table is not None
    rows = _get_report_rows(table)
    assert rows['Sick SICK-123'][0] == '8'
    assert rows['Sick SICK-123'][9] == '8'
    assert rows['Special leave (Parental leave)'][0] == '4'
    assert rows['Special leave (Parental leave)'][10] == '4'


def test_timesheet_report_marks_holidays_as_non_working_days(admin_client):
    contract = ContractFactory()
    resource = contract.resource
    DayEntryFactory(
        contract=contract,
        resource=resource,
        day=_dt('2025-06-02'),
        due_hours=Decimal(0),
        is_holiday=True,
    )

    response = admin_client.get(reverse('report-month', args=['202506']))

    assert response.status_code == 200
    table = _get_resource_table(response, resource)
    assert table is not None
    holiday_row = next(
        row
        for row in table.find_all('tr', class_='row-data')
        if row.find('td', class_='row-header').get_text(strip=True) == 'Holiday'
    )
    day_cell = holiday_row.find_all('td', class_='cell-data')[2]
    assert 'non-workday' in day_cell.get('class', [])


def test_resource_can_see_their_empty_timesheet_report_when_contractor(client):
    contractor_contract = ContractFactory(
        contract_type=Contract.ContractType.CONTRACTOR,
    )
    resource = contractor_contract.resource
    client.login(username=resource.user.username, password=resource.user._password)

    response = client.get(reverse('report-month', args=['202506']))

    assert response.status_code == 200
    table = _get_resource_table(response, resource)
    assert table is not None
    assert 'No time entries available' in table.get_text()


@pytest.fixture
def contracts_by_type():
    return {contract_type: ContractFactory(contract_type=contract_type) for contract_type in Contract.ContractType}


@pytest.mark.parametrize(
    ('query', 'expected_types'),
    [
        pytest.param('', {'EMPLOYEE'}, id='default-employee'),
        pytest.param('?contract_type=CONTRACTOR', {'CONTRACTOR'}, id='contractor-only'),
        pytest.param(
            '?contract_type=EMPLOYEE&contract_type=CONTRACTOR&contract_type=OTHER',
            {'EMPLOYEE', 'CONTRACTOR', 'OTHER'},
            id='all-types',
        ),
        pytest.param('?contract_type=INVALID', {'EMPLOYEE'}, id='invalid-falls-back-to-employee'),
    ],
)
def test_timesheet_report_contract_type_filter(admin_client, contracts_by_type, query, expected_types):
    response = admin_client.get(reverse('report-month', args=['202506']) + query)

    assert response.status_code == 200
    shown = {
        contract_type
        for contract_type, contract in contracts_by_type.items()
        if _get_resource_table(response, contract.resource) is not None
    }
    assert shown == expected_types
    soup = BeautifulSoup(response.content, 'html.parser')
    checked = {checkbox['value'] for checkbox in soup.select('input[name="contract_type"][checked]')}
    assert checked == expected_types


@pytest.mark.parametrize(
    ('permission', 'can_filter'),
    [
        pytest.param(None, False, id='no-permission'),
        pytest.param('core.view_any_timesheet', True, id='view-any'),
        pytest.param('core.manage_any_timesheet', True, id='manage-any'),
    ],
)
def test_timesheet_report_contract_type_filter_requires_permission(client, contracts_by_type, permission, can_filter):
    own_contract = ContractFactory(contract_type=Contract.ContractType.EMPLOYEE)
    user = own_contract.resource.user
    if permission:
        add_permissions(user, permission)
    client.force_login(user)

    response = client.get(reverse('report-month', args=['202506']) + '?contract_type=CONTRACTOR')

    assert response.status_code == 200
    soup = BeautifulSoup(response.content, 'html.parser')
    assert (soup.find(id='filter-section') is not None) is can_filter
    contractor = contracts_by_type[Contract.ContractType.CONTRACTOR].resource
    assert (_get_resource_table(response, contractor) is not None) is can_filter
    assert (_get_resource_table(response, own_contract.resource) is not None) is not can_filter


def test_timesheet_report_month_navigation_keeps_contract_type_filter(admin_client):
    response = admin_client.get(
        reverse('report-month', args=['202506']) + '?contract_type=CONTRACTOR&contract_type=OTHER'
    )

    soup = BeautifulSoup(response.content, 'html.parser')
    hrefs = {a['href'] for a in soup.find_all('a', class_='button')}
    assert reverse('report-month', args=['202505']) + '?contract_type=CONTRACTOR&contract_type=OTHER' in hrefs
    assert reverse('report-month', args=['202507']) + '?contract_type=CONTRACTOR&contract_type=OTHER' in hrefs


def _get_exported_resources(response, resources):
    sheet = openpyxl.load_workbook(BytesIO(response.content)).active
    headers = {str(cell.value) for cell in sheet['A'] if cell.value}
    return {
        key
        for key, resource in resources.items()
        if any(f'{resource.last_name.upper()} {resource.first_name}' in header for header in headers)
    }


@pytest.mark.parametrize(
    ('query', 'expected_types'),
    [
        pytest.param('', {'EMPLOYEE'}, id='default-employee'),
        pytest.param('?contract_type=CONTRACTOR&contract_type=OTHER', {'CONTRACTOR', 'OTHER'}, id='non-employees'),
        pytest.param('?contract_type=INVALID', {'EMPLOYEE'}, id='invalid-falls-back-to-employee'),
    ],
)
def test_timesheet_report_export_contract_type_filter(admin_client, contracts_by_type, query, expected_types):
    response = admin_client.get(reverse('export_report', args=['202506']) + query)

    assert response.status_code == 200
    resources = {contract_type: contract.resource for contract_type, contract in contracts_by_type.items()}
    assert _get_exported_resources(response, resources) == expected_types


@pytest.mark.parametrize(
    ('permission', 'can_filter'),
    [
        pytest.param(None, False, id='no-permission'),
        pytest.param('core.view_any_timesheet', True, id='view-any'),
        pytest.param('core.manage_any_timesheet', True, id='manage-any'),
    ],
)
def test_timesheet_report_export_contract_type_filter_requires_permission(
    client, contracts_by_type, permission, can_filter
):
    own_contract = ContractFactory(contract_type=Contract.ContractType.EMPLOYEE)
    user = own_contract.resource.user
    if permission:
        add_permissions(user, permission)
    client.force_login(user)

    response = client.get(reverse('export_report', args=['202506']) + '?contract_type=CONTRACTOR')

    assert response.status_code == 200
    exported = _get_exported_resources(
        response, {'own': own_contract.resource, 'contractor': contracts_by_type['CONTRACTOR'].resource}
    )
    assert exported == ({'contractor'} if can_filter else {'own'})


def test_timesheet_report_download_link_keeps_contract_type_filter(admin_client):
    response = admin_client.get(
        reverse('report-month', args=['202506']) + '?contract_type=CONTRACTOR&contract_type=OTHER'
    )

    soup = BeautifulSoup(response.content, 'html.parser')
    hrefs = {a['href'] for a in soup.find_all('a', class_='button')}
    assert reverse('export_report', args=['202506']) + '?contract_type=CONTRACTOR&contract_type=OTHER' in hrefs


@pytest.mark.parametrize(
    ('holiday_day', 'expected_days'),
    [
        pytest.param(Contract.Weekday.SUNDAY, {1, 8, 15, 22, 29}, id='sunday'),
        pytest.param(Contract.Weekday.WEDNESDAY, {4, 11, 18, 25}, id='wednesday'),
        pytest.param(None, set(), id='none'),
    ],
)
def test_timesheet_report_marks_the_weekly_holiday(admin_client, holiday_day, expected_days):
    # June 2025 in GB-ENG has no public holidays
    contract = ContractFactory(
        holiday_day=holiday_day,
        base_in__country__name='United Kingdom',
        base_in__country__country_calendar_code='GB',
        base_in__subdivision_code='ENG',
    )

    response = admin_client.get(reverse('report-month', args=['202506']))
    table = _get_resource_table(response, contract.resource)
    header_cells = table.find('thead').find_all('tr')[0].find_all('td')[1:]
    assert {day for day, cell in enumerate(header_cells, 1) if cell.get_text(strip=True) == 'X'} == expected_days

    response = admin_client.get(reverse('export_report', args=['202506']))
    sheet = openpyxl.load_workbook(BytesIO(response.content)).active
    header_row = next(row for row in sheet.iter_rows(values_only=True) if contract.resource.first_name in str(row[0]))
    assert {day for day, value in enumerate(header_row[2:], 1) if value == 'X'} == expected_days
